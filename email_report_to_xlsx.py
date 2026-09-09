#!/usr/bin/env python3
"""Rebuild a Multibrand Flash Report .xlsx from the emailed attachment's text.

The report is emailed every morning from reports@report-emailer.com carrying the
prior business day. The Outlook connector hands the .xlsx attachment back as
tab-delimited text rather than the original bytes, so the workbook is rebuilt
here cell for cell.

That rebuild is faithful, not approximate: rebuilding 2026-09-06 from its own
email and parsing both copies with data_loader produced identical payloads
across all eight stores and all four sections. The text extraction drops blank
rows, which shifts row numbers -- immaterial, because data_loader anchors every
section on its header text rather than on fixed rows.

  python3 email_report_to_xlsx.py report.tsv            # name it by its own date
  python3 email_report_to_xlsx.py report.tsv out.xlsx   # or name it yourself

Exit codes: 0 written, 1 the text did not check out.
"""
import datetime
import os
import re
import sys

from openpyxl import Workbook, load_workbook

SHEET = 'Multibrand_DailyFlashReport'
MAX_COL = 21
PREFIX = 'Multibrand_FlashReport'
_NUMBER = re.compile(r'^-?\d+(\.\d+)?$')
_DATE = re.compile(r'Selected Date:\s*(\d{1,2})/(\d{1,2})/(\d{4})')
_STORE = re.compile(r'^\d{4}\s*-\s*')


def cell_value(text):
    """A number where the source held a number, a string where it held text."""
    if text == '':
        return None
    if _NUMBER.match(text):
        return int(text) if '.' not in text else float(text)
    return text


def build(tsv_text, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    row = 0
    for line in tsv_text.split('\n'):
        line = line.rstrip('\r')
        if line.startswith('=== Sheet:') or line.strip() == '':
            continue                      # connector banner, and dropped blanks
        row += 1
        for col, raw in enumerate(line.split('\t')[:MAX_COL], start=1):
            value = cell_value(raw)
            if value is not None:
                ws.cell(row=row, column=col, value=value)
    wb.save(out_path)
    return row


def report_date(tsv_text):
    match = _DATE.search(tsv_text)
    if not match:
        return None
    mm, dd, yyyy = (int(g) for g in match.groups())
    return datetime.date(yyyy, mm, dd)


SECTION_KEYS = {
    'day sales': 'sales',
    'day trans': 'transactions',
    'average check': 'channels',
    'labor': 'labor',
}


def _sections(grid):
    """Per-store row runs keyed by section, the way data_loader finds them.

    Mirroring data_loader's own anchoring is the point: a workbook that passes
    here is one the dashboard can actually read.
    """
    runs, current = [], []
    for cells in grid:
        label = cells[0]
        if label is not None and _STORE.match(str(label).strip()):
            current.append(cells)
        elif current:
            runs.append(current)
            current = []
    if current:
        runs.append(current)

    found = {}
    for index, run in enumerate(runs):
        start = grid.index(run[0])
        name = None
        for row in range(start - 1, max(-1, start - 6), -1):
            joined = ' '.join(str(c or '') for c in grid[row]).lower()
            for key, section in SECTION_KEYS.items():
                if key in joined:
                    name = section
                    break
            if name:
                break
        if name and name not in found:
            found[name] = run
    return found


def verify(path, day):
    """Refuse to hand back a workbook that does not hold a whole day.

    Checks the date survived, that all four per-store sections are present with
    eight stores each, and that per-store YTD sales add up to the brand total.
    A truncated paste keeps the sales block and loses the rest, so the section
    check -- not the arithmetic -- is what actually catches it.
    """
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        if SHEET not in wb.sheetnames:
            return 'sheet {} is missing'.format(SHEET)
        ws = wb[SHEET]
        grid = [[c.value for c in r] for r in
                ws.iter_rows(min_row=1, max_row=120, max_col=MAX_COL)]
    finally:
        wb.close()

    header = ' '.join(str(c) for c in grid[1] if c is not None)
    if not _DATE.search(header):
        return 'no readable date in cell A2'
    if report_date(header) != day:
        return 'date in the workbook is not {}'.format(day)

    stores, brand = {}, None
    for cells in grid:
        label = cells[0]
        if label is None:
            continue
        text = str(label).strip()
        if _STORE.match(text):
            stores.setdefault(text, cells[13])          # column N = YTD sales TY
        elif text.startswith('Brand Totals') and brand is None:
            brand = cells[13]

    if len(stores) != 8:
        return 'found {} stores, expected 8'.format(len(stores))
    if brand is None:
        return 'no Brand Totals row'
    total = sum(v for v in stores.values() if isinstance(v, (int, float)))
    if abs(total - brand) > 0.01:
        return 'stores sum to {:,.2f} but brand total is {:,.2f}'.format(total, brand)

    sections = _sections(grid)
    for name in sorted(SECTION_KEYS.values()):
        if name not in sections:
            return 'no {} section -- the text looks truncated'.format(name)
        if len(sections[name]) != 8:
            return '{} section has {} stores, expected 8'.format(name, len(sections[name]))
    return None


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 1
    with open(sys.argv[1], encoding='utf-8') as fh:
        text = fh.read()

    day = report_date(text)
    if day is None:
        print('!! no "Selected Date:" line in the text -- refusing to write')
        return 1

    out = sys.argv[2] if len(sys.argv) > 2 else '{}.{}.xlsx'.format(PREFIX, day.isoformat())
    rows = build(text, out)

    problem = verify(out, day)
    if problem:
        os.remove(out)
        print('!! rebuilt workbook failed its check ({}) -- not written'.format(problem))
        return 1

    print('wrote {}  ({} rows, business date {}, 8 stores, totals tie)'
          .format(out, rows, day.isoformat()))
    return 0


if __name__ == '__main__':
    sys.exit(main())
